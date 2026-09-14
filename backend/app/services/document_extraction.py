"""Plain-text extraction from uploaded CV files. No AI/LLM involved — pypdf/python-docx do
structural parsing only (spec Rule 8: never claim AI where there is none)."""

import io
import zipfile

import docx
from fastapi import HTTPException, status
from pypdf import PdfReader

# Phase 11 (spec §98) resource limits. The 5MB upload cap alone doesn't bound CPU/memory: a small
# PDF can hold thousands of pages, and a DOCX is a zip that can decompress to gigabytes.
MAX_PDF_PAGES = 20  # far beyond any real CV; anything longer is rejected, not truncated silently
MAX_DOCX_UNCOMPRESSED_BYTES = 30 * 1024 * 1024
MAX_EXTRACTED_CHARS = 100_000  # bounds the ATS engine's work and the stored cv_documents row

SUPPORTED_CONTENT_TYPES = {
    "application/pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",  # .docx
    "text/plain",
}


def extract_text(*, filename: str | None, content: bytes, content_type: str | None) -> str:
    ext = (filename or "").lower().rsplit(".", 1)[-1] if filename and "." in filename else ""

    if content_type == "application/pdf" or ext == "pdf":
        text = _extract_pdf_text(content)
    elif (
        content_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        or ext == "docx"
    ):
        text = _extract_docx_text(content)
    elif content_type == "text/plain" or ext == "txt":
        text = content.decode("utf-8", errors="ignore")
    else:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="Unsupported CV format. Upload a PDF, DOCX, or plain-text file.",
        )

    if len(text) > MAX_EXTRACTED_CHARS:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="This CV contains far more text than a CV normally would. Please upload a shorter version.",
        )
    return text


def _extract_pdf_text(content: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(content))
        page_count = len(reader.pages)
    except Exception as exc:  # pypdf raises various exception types for malformed PDFs
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="Could not read this PDF file."
        ) from exc

    if page_count > MAX_PDF_PAGES:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=f"This PDF has {page_count} pages. CVs longer than {MAX_PDF_PAGES} pages aren't supported.",
        )

    try:
        pages = [page.extract_text() or "" for page in reader.pages]
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="Could not read this PDF file."
        ) from exc

    text = "\n".join(pages).strip()
    if not text:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="No extractable text found in this PDF (it may be a scanned image without OCR).",
        )
    return text


def _extract_docx_text(content: bytes) -> str:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            uncompressed = sum(info.file_size for info in archive.infolist())
    except zipfile.BadZipFile as exc:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="Could not read this DOCX file."
        ) from exc
    if uncompressed > MAX_DOCX_UNCOMPRESSED_BYTES:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="This DOCX file expands to an unexpectedly large size and can't be processed.",
        )

    try:
        document = docx.Document(io.BytesIO(content))
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="Could not read this DOCX file."
        ) from exc

    paragraphs = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.append(cell.text)

    text = "\n".join(p for p in paragraphs if p.strip())
    if not text.strip():
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="No extractable text found in this DOCX file."
        )
    return text
